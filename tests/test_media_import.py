import asyncio
from io import BytesIO
from email.message import EmailMessage
import httpx
from pptx import Presentation
from intel_sourcerer.media_import import import_file,import_podcast_feed,save_item
from intel_sourcerer import github_import,discovery
from intel_sourcerer.storage import connect_database


def test_audio_video_and_newsletter(tmp_path):
    dbpath=tmp_path/'articles.db'
    for name,kind in [('talk.mp3','audio'),('talk.webm','video')]:
        url=asyncio.run(import_file(dbpath,b'example '+name.encode(),name))
        db=connect_database(dbpath);row=db.execute('SELECT * FROM articles WHERE url=?',(url,)).fetchone()
        assert row['content_format']==kind and '/local-media/' in row['content_markdown'];db.close()
    message=EmailMessage();message['Subject']='Research newsletter';message['From']='Researcher <test@example.com>';message.set_content('Useful intelligence.');message.add_alternative('<article><h1>Research</h1><p>Useful intelligence.</p></article>',subtype='html')
    url=asyncio.run(import_file(dbpath,message.as_bytes(),'newsletter.eml'))
    db=connect_database(dbpath);row=db.execute('SELECT * FROM articles WHERE url=?',(url,)).fetchone();assert row['title']=='Research newsletter' and row['content_format']=='newsletter' and 'Useful intelligence' in row['content_markdown'];db.close()


def test_pptx_import(tmp_path):
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text='Threat analysis';slide.placeholders[1].text='Important finding'
    stream=BytesIO();p.save(stream);dbpath=tmp_path/'articles.db';url=asyncio.run(import_file(dbpath,stream.getvalue(),'talk.pptx'))
    db=connect_database(dbpath);row=db.execute('SELECT * FROM articles WHERE url=?',(url,)).fetchone();assert row['content_format']=='slides' and 'Slide 1' in row['content_markdown'] and 'Important finding' in row['content_markdown'];db.close()


def test_podcast_feed(tmp_path,monkeypatch):
    monkeypatch.setattr(discovery,'public_url',lambda url:url)
    feed=b'<rss version="2.0"><channel><title>Security podcast</title><item><title>Latest episode</title><description>Episode notes</description><enclosure url="https://example.com/episode.mp3" type="audio/mpeg" length="10"/></item></channel></rss>'
    url=asyncio.run(import_podcast_feed(tmp_path/'articles.db',feed,'https://example.com/feed.xml'))
    assert url=='https://example.com/episode.mp3'


def test_github_readme_import(tmp_path,monkeypatch):
    import base64
    original=httpx.AsyncClient
    transport=httpx.MockTransport(lambda request:httpx.Response(200,json={'content':base64.b64encode(b'# Research tool\nUseful findings').decode()}))
    monkeypatch.setattr(github_import.httpx,'AsyncClient',lambda **kw:original(transport=transport,**kw))
    url=asyncio.run(github_import.import_github(tmp_path/'articles.db','https://github.com/example/research'))
    db=connect_database(tmp_path/'articles.db');row=db.execute('SELECT * FROM articles WHERE url=?',(url,)).fetchone();assert row['content_format']=='repository' and 'Useful findings' in row['content_markdown'];db.close()


def test_upload_reader_and_range_playback(tmp_path):
    import threading,wave
    from bs4 import BeautifulSoup
    from intel_sourcerer.web import create_server
    path=tmp_path/'articles.db';connect_database(path).close()
    buffer=BytesIO()
    with wave.open(buffer,'wb') as wav:
        wav.setnchannels(1);wav.setsampwidth(2);wav.setframerate(8000);wav.writeframes(b'\0\0'*800)
    payload=buffer.getvalue()
    server=create_server(path,0, auth_enabled=False);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    try:
        with httpx.Client(base_url=f'http://127.0.0.1:{server.server_port}',trust_env=False) as client:
            token=BeautifulSoup(client.get('/').text,'html.parser').select_one('meta[name="csrf-token"]')['content']
            response=client.post('/api/file-upload',content=payload,headers={'X-CSRF-Token':token,'X-Filename':'episode.wav'})
            assert response.status_code==200
            article=client.get('/api/article',params={'url':response.json()['url']}).json()
            assert article['content_format']=='audio'
            soup=BeautifulSoup(article['body_html'],'html.parser');block=soup.select_one('[data-media-kind="audio"]');assert block
            asset=block['data-media-url'];part=client.get(asset,headers={'Range':'bytes=0-9'})
            assert part.status_code==206 and part.content==payload[:10]
            assert client.get('/local-media/../../etc/passwd').status_code==404
            assert client.post('/api/file-upload',content=payload,headers={'X-Filename':'episode.wav'}).status_code==403
    finally:
        server.shutdown();server.server_close()
