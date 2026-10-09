"""Import audio, newsletters, slide decks, and local video into the library."""
from pathlib import Path
from io import BytesIO
from html import escape
from email import policy
from email.parser import BytesParser
from datetime import datetime, timezone
from urllib.parse import urlsplit, unquote
import hashlib
import feedparser
from bs4 import BeautifulSoup
from .content import extract_markdown, render_markdown
from .config import PublisherSource
from .feeds import Article
from .storage import connect_database, save_source, save_article

AUDIO = {'.mp3', '.m4a', '.ogg', '.wav'}
VIDEO = {'.mp4', '.webm'}
DOCUMENTS = {'.eml', '.pptx'}
SUPPORTED = AUDIO | VIDEO | DOCUMENTS | {'.pdf'}


def save_item(database, url, title, markdown, content_format, publisher, author=None):
    db = connect_database(database)
    try:
        existing = db.execute('SELECT 1 FROM articles WHERE url=?', (url,)).fetchone()
        if existing:
            return url
        source = PublisherSource('media-' + hashlib.sha256(publisher.encode()).hexdigest()[:12], publisher, 'https://example.invalid', 'manual')
        now = datetime.now(timezone.utc)
        save_source(db, source)
        text = BeautifulSoup(render_markdown(markdown), 'html.parser').get_text(' ', strip=True)
        save_article(db, source.id, Article(publisher,title,url,url,now,author,None,''), content=text, error=None, reason='Manually imported media; publication date defaults to import time', discovered_at=now.isoformat(),markdown=markdown)
        db.execute('UPDATE articles SET content_format=? WHERE url=?', (content_format,url))
        db.commit()
        return url
    finally:
        db.close()


def media_markdown(title, url, kind):
    # Generate via HTML so titles cannot inject Markdown links or formatting.
    return extract_markdown(f'<article><blockquote><p><strong>{kind.capitalize()} placeholder: </strong>{escape(title)}</p><p><a href="{escape(url,quote=True)}">Open original {kind}</a></p></blockquote></article>',url)[0]


def store_asset(database, payload, suffix):
    directory = Path(database).resolve().parent / 'media'
    directory.mkdir(parents=True, exist_ok=True)
    name = hashlib.sha256(payload).hexdigest() + suffix.lower()
    path = directory / name
    if not path.exists():
        path.write_bytes(payload)
    return '/local-media/' + name


async def import_file(database, payload, filename):
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError('Supported files: PDF, MP3, M4A, OGG, WAV, MP4, WebM, EML and PPTX.')
    limit = 20_000_000 if suffix in DOCUMENTS | {'.pdf'} else 100_000_000
    if not payload or len(payload) > limit:
        raise ValueError(f'This file exceeds the {limit // 1_000_000} MB limit or is empty.')
    if suffix == '.pdf':
        from .discovery import import_pdf_upload
        return await import_pdf_upload(database,payload)
    title = Path(filename).stem
    url = store_asset(database,payload,suffix)
    # Stable identity separate from the locally served media URL.
    identity = (Path(database).resolve().parent / 'media' / url.rsplit('/',1)[-1]).as_uri()
    if suffix in AUDIO | VIDEO:
        kind = 'audio' if suffix in AUDIO else 'video'
        return save_item(database,identity,title,media_markdown(title,url,kind),kind,'Local media')
    if suffix == '.eml':
        message = BytesParser(policy=policy.default).parsebytes(payload)
        title = str(message.get('Subject') or title)
        body = message.get_body(preferencelist=('html','plain'))
        if body is None:
            raise ValueError('This email has no readable body.')
        value = body.get_content()
        html = value if body.get_content_type() == 'text/html' else '<article><p>' + escape(value).replace('\n','<br>') + '</p></article>'
        markdown,_ = extract_markdown(html,'https://newsletter.invalid/')
        return save_item(database,identity,title,markdown,'newsletter','Email newsletters',str(message.get('From') or ''))
    from pptx import Presentation
    import zipfile
    if not zipfile.is_zipfile(BytesIO(payload)):
        raise ValueError('This file is not a valid PPTX slide deck.')
    with zipfile.ZipFile(BytesIO(payload)) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 100_000_000:
            raise ValueError('This slide deck expands beyond the 100 MB limit.')
    presentation = Presentation(BytesIO(payload))
    if len(presentation.slides) > 300:
        raise ValueError('Slide decks must contain no more than 300 slides.')
    html = ['<article>']
    for index,slide in enumerate(presentation.slides,1):
        html.append(f'<h2>Slide {index}</h2>')
        for shape in slide.shapes:
            if shape.has_text_frame:
                html.append('<p>'+escape(shape.text).replace('\n','<br>')+'</p>')
            if shape.has_table:
                html.append('<table>' + ''.join('<tr>'+''.join('<td>'+escape(cell.text)+'</td>' for cell in row.cells)+'</tr>' for row in shape.table.rows)+'</table>')
            if hasattr(shape,'image'):
                image = shape.image
                # Convert to PNG so figures use a predictable browser-safe format.
                from PIL import Image
                output = BytesIO()
                Image.open(BytesIO(image.blob)).convert('RGB').save(output,format='PNG')
                src = store_asset(database,output.getvalue(),'.png')
                html.append(f'<img src="{src}" alt="Slide {index} figure"/>')
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                html.append('<h3>Speaker notes</h3><p>'+escape(notes)+'</p>')
    html.append('</article>')
    markdown,_ = extract_markdown(''.join(html),'https://slides.invalid/')
    return save_item(database,identity,title,markdown,'slides','Slide decks')


async def import_podcast_feed(database, payload, url):
    from .discovery import public_url
    import asyncio
    feed = feedparser.parse(payload)
    for entry in feed.entries:
        enclosure = next((item for item in entry.get('enclosures',[]) if item.get('type','').startswith('audio/') or Path(urlsplit(item.get('href','')).path).suffix.lower() in AUDIO),None)
        if not enclosure:
            continue
        audio = await asyncio.to_thread(public_url,enclosure['href'])
        title = entry.get('title','Podcast episode')
        summary = entry.get('summary','')
        description,_ = extract_markdown('<article>'+summary+'</article>',url)
        return save_item(database,audio,title,media_markdown(title,audio,'audio')+'\n\n'+description,'audio',feed.feed.get('title','Podcasts'))
    raise ValueError('No audio episodes found in this podcast feed.')
